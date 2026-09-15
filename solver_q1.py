# -*- coding: utf-8 -*-
"""
2026 高教社杯 A题 问题1 求解器
药材热风烘干 · 预热平衡阶段 (0-1800 s)

模型：一维径向圆柱（忽略端部效应）
  传热: rho*cp*dT/dt = (1/r)d/dr(r k dT/dr)
  传质: dC/dt = (1/r)d/dr(r D(C) dC/dr),  D(C)=7e-9*exp(-0.89/C)
边界: r=0 对称; r=R 第三类(对流传热 h, 对流传质 hm)
初值: T=28 C, C=2.55 kg/kg
环境: 附件1 T_a(t), C_a(t) 线性插值

数值: 空间有限体积(FV) + 时间全隐式向后欧拉 + Picard 迭代(处理 D(C) 非线性)
输出: result1.xlsx (温度/水分浓度 两工作表, 1s x 0.1cm, 四位小数)
"""
import numpy as np
import openpyxl
from pathlib import Path

# ============ 参数 ============
R   = 0.02        # 半径 m
rho = 820.0       # 湿药材表观密度 kg/m^3 (附录2)
cp  = 2600.0      # 比热 J/(kg K)
k   = 0.36        # 导热 W/(m K)
h   = 25.0        # 对流换热 W/(m^2 K)
hm  = 8e-7        # 对流传质 m/s
T0  = 28.0
C0  = 2.55

def Dfun(C):
    """水分扩散系数 (m^2/s), 附录2"""
    return 7e-9*np.exp(-0.89/C)

# ============ 附件1 环境 ============
def load_ambient(xlsx):
    wb = openpyxl.load_workbook(xlsx, data_only=True)
    rows = list(wb.active.iter_rows(values_only=True))[1:]
    t  = np.array([r[0] for r in rows], float)
    Ta = np.array([r[1] for r in rows], float)
    Ca = np.array([r[2] for r in rows], float)
    return (lambda tt: np.interp(tt, t, Ta)), (lambda tt: np.interp(tt, t, Ca))

# ============ FV 网格 ============
def build_grid(N):
    dr = R/N
    r  = np.arange(N+1)*dr
    rf = (np.arange(N)+0.5)*dr          # 内部面半径(面 i 介于节点 i 与 i+1)
    Af = 2*np.pi*rf                     # 面面积(略去长度 L, 处处约去)
    AR = 2*np.pi*R                      # 表面面积
    V  = np.zeros(N+1)
    V[0] = np.pi*(dr/2)**2
    V[1:-1] = np.pi*(rf[1:]**2 - rf[:-1]**2)
    V[-1] = np.pi*(R**2 - rf[-1]**2)
    return dr, r, rf, Af, AR, V

# ============ 三对角求解(Thomas) ============
def thomas(a, b, c, d):
    n = len(d)
    cp = np.zeros(n); dp = np.zeros(n)
    cp[0] = c[0]/b[0]; dp[0] = d[0]/b[0]
    for i in range(1, n):
        m = b[i] - a[i]*cp[i-1]
        cp[i] = c[i]/m
        dp[i] = (d[i] - a[i]*dp[i-1])/m
    x = np.zeros(n); x[-1] = dp[-1]
    for i in range(n-2, -1, -1):
        x[i] = dp[i] - cp[i]*x[i+1]
    return x

# ============ 单次求解 ============
def solve(N, dt, t_end, Tamb, Camb, out_times):
    dr, r, rf, Af, AR, V = build_grid(N)
    r_out = np.linspace(0, R, 21)        # 0,0.1,...,2.0 cm
    T = np.full(N+1, T0)
    C = np.full(N+1, C0)

    # 传热系数矩阵(常数, 仅 Q1)
    aT = np.zeros(N+1); bT = np.zeros(N+1); cT = np.zeros(N+1)
    bT[1:-1] = rho*cp*V[1:-1]/dt + (k/dr)*(Af[:-1]+Af[1:])
    aT[1:-1] = -(k/dr)*Af[:-1]
    cT[1:-1] = -(k/dr)*Af[1:]
    bT[0] = rho*cp*V[0]/dt + (k/dr)*Af[0]
    cT[0] = -(k/dr)*Af[0]
    bT[-1] = rho*cp*V[-1]/dt + (k/dr)*Af[-1] + h*AR
    aT[-1] = -(k/dr)*Af[-1]

    out = {t: None for t in out_times}
    t = 0.0
    flux_int = 0.0          # 表面归一化水分通量时间积分 (m^3)
    step = 0
    while t < t_end - 1e-12:
        tnext = min(t+dt, t_end)
        dtn = tnext - t
        # ---- 传热: 隐式线性 ----
        d = np.zeros(N+1)
        d[1:-1] = rho*cp*V[1:-1]/dt*T[1:-1]
        d[0]   = rho*cp*V[0]/dt*T[0]
        d[-1]  = rho*cp*V[-1]/dt*T[-1] + h*AR*Tamb(tnext)
        Tnew = thomas(aT, bT, cT, d)

        # ---- 传质: 隐式 + Picard ----
        Cguess = C.copy()
        Ca_n = Camb(tnext)
        for it in range(30):
            Dn = Dfun(Cguess)
            Df = 2*Dn[:-1]*Dn[1:]/(Dn[:-1]+Dn[1:]+1e-300)
            a = np.zeros(N+1); b = np.zeros(N+1); c = np.zeros(N+1); dd = np.zeros(N+1)
            b[1:-1] = V[1:-1]/dt + (Af[:-1]*Df[:-1] + Af[1:]*Df[1:])/dr
            a[1:-1] = -(Af[:-1]*Df[:-1])/dr
            c[1:-1] = -(Af[1:]*Df[1:])/dr
            b[0] = V[0]/dt + Af[0]*Df[0]/dr
            c[0] = -(Af[0]*Df[0])/dr
            b[-1] = V[-1]/dt + Af[-1]*Df[-1]/dr + hm*AR
            a[-1] = -(Af[-1]*Df[-1])/dr
            dd[1:-1] = V[1:-1]/dt*C[1:-1]
            dd[0]  = V[0]/dt*C[0]
            dd[-1] = V[-1]/dt*C[-1] + hm*AR*Ca_n
            Cnew = thomas(a, b, c, dd)
            if np.max(np.abs(Cnew-Cguess)) < 1e-12:
                Cguess = Cnew
                break
            Cguess = Cnew

        # 通量积分(梯形, 归一化, m^3)
        flux_old = hm*(C[-1]-Camb(t))
        flux_new = hm*(Cnew[-1]-Ca_n)
        flux_int += 0.5*(flux_old+flux_new)*AR*dtn

        T = Tnew; C = Cnew; t = tnext; step += 1

        if t in out:
            out[t] = (np.interp(r_out, r, T), np.interp(r_out, r, C))

    return out, T, C, flux_int, V

# ============ 主流程 ============
if __name__ == "__main__":
    base = Path(__file__).resolve().parent
    xlsx1 = base/"附件"/"附件1.xlsx"
    Tamb, Camb = load_ambient(xlsx1)

    table_t = [100,300,600,900,1200,1500,1800]
    table_r = [0,0.005,0.01,0.015,0.02]

    # ---- 基线: N=400, dt=1s ----
    N = 400; dt = 1.0; t_end = 1800.0
    out_times = np.arange(1, 1801)      # 每秒
    out, Tf, Cf, flux_int, V = solve(N, dt, t_end, Tamb, Camb, out_times)

    print("="*70)
    print(f"Q1 求解完成  N={N} (dr={R/N*100:.3f} cm), dt={dt} s, t_end={t_end} s")
    print("="*70)

    # 表1 温度 / 表2 水分浓度
    idx = [int(round(rr/(R/N))) for rr in table_r]
    print("\n表1 温度 (℃)")
    print("t/s     r=0     0.5     1.0     1.5     2.0")
    for tt in table_t:
        row = out[tt][0]
        print(f"{tt:5d} " + " ".join(f"{row[int(round(rr/0.001))]:8.4f}" for rr in table_r))
    print("\n表2 水分浓度 (kg/kg)")
    print("t/s     r=0     0.5     1.0     1.5     2.0")
    for tt in table_t:
        row = out[tt][1]
        print(f"{tt:5d} " + " ".join(f"{row[int(round(rr/0.001))]:8.4f}" for rr in table_r))

    # ---- 检验 ----
    print("\n" + "="*70)
    print("模型检验")
    print("="*70)
    # 对称性
    print(f"对称性: dT/dr@0 ~ { (Tf[1]-Tf[0])/(R/N):.3e} ℃/m ; dC/dr@0 ~ {(Cf[1]-Cf[0])/(R/N):.3e} (kg/kg)/m")
    # 表面边界残差
    resT = k*(Tf[-1]-Tf[-2])/(R/N) - h*(Tamb(1800)-Tf[-1])
    resC = Dfun(Cf[-1])*(Cf[-1]-Cf[-2])/(R/N) - hm*(Camb(1800)-Cf[-1])
    print(f"表面边界残差: 热 {resT:.3e} ; 质 {resC:.3e}  (应≈0)")
    # 质量守恒(归一化, m^3)
    loss = np.sum((C0-Cf)*V)
    print(f"水分守恒: 域内减少 {loss:.6e} m^3 ; 表面通量积分 {flux_int:.6e} m^3 ; 相对误差 {abs(loss-flux_int)/max(loss,1e-30)*100:.4f} %")

    # ---- 网格无关性 ----
    print("\n网格/时间步无关性(表1 温度 @1800s 表面 & 中心):")
    for NN in [200,400,800]:
        o,_,_,_,_ = solve(NN, dt, t_end, Tamb, Camb, [1800.0])
        row = o[1800.0][0]
        print(f"  N={NN:4d} (dr={R/NN*100:.4f} cm): T_center={row[0]:.4f}, T_surf={row[-1]:.4f}")
    print("时间步无关性(N=400, 表1 @1800s):")
    for dtt in [2.0,1.0,0.5]:
        o,_,_,_,_ = solve(N, dtt, t_end, Tamb, Camb, [1800.0])
        row = o[1800.0][0]
        print(f"  dt={dtt:4.1f} s: T_center={row[0]:.4f}, T_surf={row[-1]:.4f}")

    # ---- 写 result1.xlsx ----
    outdir = base/"outputs"
    outdir.mkdir(exist_ok=True)
    r_vals = np.linspace(0, R, 21)
    times = np.arange(1, 1801)
    wb = openpyxl.Workbook()
    ws1 = wb.active; ws1.title = "温度"
    ws2 = wb.create_sheet("水分浓度")
    header = ["时间\\到药材中心的距离"] + [round(float(rr)*100,1) for rr in r_vals]
    ws1.append(header); ws2.append(header)
    for k, tt in enumerate(times, start=2):
        Tr, Cr = out[tt]
        ws1.cell(k,1,float(tt))
        ws2.cell(k,1,float(tt))
        for j in range(21):
            ws1.cell(k,j+2,round(float(Tr[j]),4))
            ws2.cell(k,j+2,round(float(Cr[j]),4))
    for ws in (ws1, ws2):
        for row in ws.iter_rows(min_row=2, min_col=2, max_row=1801, max_col=22):
            for cell in row:
                cell.number_format = '0.0000'
    res_path = outdir/"result1.xlsx"
    wb.save(res_path)
    print(f"\n已写出: {res_path}")
    print(f"  尺寸: {ws1.max_row-1} 行 x {ws1.max_column-1} 列")
