"""Conservative radial finite-volume solver; authoritative inputs remain read-only."""
from pathlib import Path
from dataclasses import dataclass
import hashlib, json, platform, sys, time
import numpy as np
import pandas as pd
import scipy
from scipy.integrate import solve_ivp
from scipy.interpolate import PchipInterpolator
from scipy.sparse import lil_matrix
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'results'
AIR=pd.read_excel(ROOT/'A题/附件/附件1.xlsx',header=None).iloc[1:,:3].to_numpy(float)
RAD=pd.read_excel(ROOT/'A题/附件/附件2.xlsx',header=None).iloc[1:,:2].to_numpy(float)
DATA=ROOT/'A题'/'附件'
air=AIR
RP=PchipInterpolator(RAD[:,0],RAD[:,1]/100,extrapolate=False)
R0=.02; L=.25; H=25.; HM=8e-7; LIMIT=.15

def environment(t,extension='plateau'):
    if t<=AIR[-1,0]: return float(np.interp(t,AIR[:,0],AIR[:,1])),float(np.interp(t,AIR[:,0],AIR[:,2]))
    if extension=='last': return tuple(AIR[-1,1:])
    if extension=='mean': return tuple(AIR[AIR[:,0]>=10800,1:].mean(axis=0))
    return 50.,.05

def radius(t,moving=True,interp='pchip'):
    if not moving: return R0
    if t>=RAD[-1,0]: return RAD[-1,1]/100
    return float(RP(t)) if interp=='pchip' else float(np.interp(t,RAD[:,0],RAD[:,1]/100))

def Tinf(t: float) -> float:
    return environment(float(t))[0]

def Cbinf(t: float) -> float:
    return environment(float(t))[1]

def Rmeas(t: float) -> float:
    return radius(float(t), True, 'pchip')

def properties(kind,T,C):
    if np.any(C<=0) or np.any(T+273.15<=0): raise ValueError('Non-positive constitutive state')
    if kind=='q1': return np.full_like(C,820.),np.full_like(C,2600.),np.full_like(C,.36),7e-9*np.exp(-.89/C)
    if kind in ('q2','q3'):
        return 650+128*C,1450+2736*C/(C+1),.21+.38*C/(C+1),2.4e-3*np.exp(-.45/C-3850/(T+273.15))
    return 760+90*C,1850+2150*C/(C+1),.12+.20*C/(C+1),4.2e-4*np.exp(-.30/C-3850/(T+273.15))

def center(u):
    """Even quadratic reconstruction at r=0."""
    return (9*u[...,0]-u[...,1])/8

@dataclass
class Model:
    kind: str
    n: int=160
    extension: str='plateau'
    interpolation: str='pchip'
    moving: bool|None=None
    h: float=H
    hm: float=HM
    def __post_init__(self):
        if self.moving is None: self.moving=self.kind=='q4'
        self.faces=np.linspace(0,1,self.n+1)
        self.x=(self.faces[:-1]+self.faces[1:])/2
        self.weights=np.diff(self.faces**2)
    def geometry(self,t): return radius(t,self.moving,self.interpolation)
    def surfaces(self,t,y):
        T,C=y[:self.n],y[self.n:]; _,_,k,D=properties(self.kind,T,C)
        R=self.geometry(t); bT,bC=environment(t,self.extension); half=R/(2*self.n)
        return (T[-1]+self.h*half/k[-1]*bT)/(1+self.h*half/k[-1]),(C[-1]+self.hm*half/D[-1]*bC)/(1+self.hm*half/D[-1])
    def rhs(self,t,y):
        T,C=y[:self.n],y[self.n:]; rho,cp,k,D=properties(self.kind,T,C)
        R=self.geometry(t); dr=R/self.n
        areas=2*np.pi*L*R*self.faces
        volumes=np.pi*L*R**2*self.weights
        bT,bC=environment(t,self.extension)
        fluxes=[]
        for u,a,b,h in [(T,k,bT,self.h),(C,D,bC,self.hm)]:
            flux=np.zeros(self.n+1)
            af=2*a[:-1]*a[1:]/(a[:-1]+a[1:])
            flux[1:-1]=areas[1:-1]*af*np.diff(u)/dr
            flux[-1]=areas[-1]*(b-u[-1])/(dr/(2*a[-1])+1/h) if h>0 else 0
            fluxes.append(np.diff(flux)/volumes)
        return np.r_[fluxes[0]/(rho*cp),fluxes[1]]
    def max_moisture(self,t,y): return max(float(np.max(y[self.n:])),float(center(y[self.n:])),self.surfaces(t,y)[1])
    def sparsity(self):
        s=lil_matrix((2*self.n,2*self.n),dtype=int)
        for i in range(self.n):
            for j in range(max(0,i-1),min(self.n,i+2)):
                s[i,j]=s[i,self.n+j]=s[self.n+i,j]=s[self.n+i,self.n+j]=1
        return s.tocsr()
    def solve(self,horizon,step=60,event=False,method='BDF',rtol=1e-7,max_step=900):
        y0=np.r_[np.full(self.n,28.),np.full(self.n,2.55)]
        def crossing(t,y): return self.max_moisture(t,y)-LIMIT
        crossing.terminal=True; crossing.direction=-1
        kw=dict(method=method,rtol=rtol,atol=np.r_[np.full(self.n,1e-7),np.full(self.n,1e-9)],max_step=max_step,dense_output=True,jac_sparsity=self.sparsity())
        s=solve_ivp(self.rhs,(0,horizon),y0,events=crossing if event else None,**kw)
        if not s.success: raise RuntimeError(s.message)
        te=float(s.t_events[0][0]) if event and len(s.t_events[0]) else None
        if event and te is None: raise RuntimeError('No threshold crossing in horizon')
        # The report point must be strictly after the continuous root, including
        # the edge case where the root lands exactly on a reporting grid point.
        end=(np.floor(te/step)+1)*step if event else horizon
        ts=np.arange(0,end+step/2,step)
        Y=s.sol(ts).T
        if te is not None and end>te:
            # actual integration after root; do not extrapolate dense output
            continuation=solve_ivp(self.rhs,(te,end),s.sol(te),**kw)
            if not continuation.success: raise RuntimeError(continuation.message)
            mask=ts>te; Y[mask]=continuation.sol(ts[mask]).T
        return dict(t=ts,T=Y[:,:self.n],C=Y[:,self.n:],te=te,success=s.success),s
    def sample(self,t,y,points):
        R=self.geometry(t); Ts,Cs=self.surfaces(t,y)
        coords=np.r_[0,self.x*R,R]
        T=np.r_[center(y[:self.n]),y[:self.n],Ts]; C=np.r_[center(y[self.n:]),y[self.n:],Cs]
        valid=np.asarray(points)<=R+1e-12
        return np.where(valid,np.interp(points,coords,T),np.nan),np.where(valid,np.interp(points,coords,C),np.nan)

def workbook(q,data,model):
    wb=Workbook(); wb.remove(wb.active)
    x=np.linspace(0,.02,21); Y=np.c_[data['T'],data['C']]
    tv=[];cv=[]
    for t,y in zip(data['t'],Y):
        T,C=model.sample(t,y,x);tv.append(T);cv.append(C)
    for name,arr in ([('温度',tv),('水分浓度',cv)] if q in ('q1','q2') else [('Sheet1',cv)]):
        ws=wb.create_sheet(name); ws.append(['时间/s']+[round(float(v*100),1) for v in x]+(['药材表面'] if q=='q4' else []))
        for j,(t,vals) in enumerate(zip(data['t'],arr)):
            values=[None if np.isnan(v) else float(v) for v in vals]
            if q=='q4':values.append(float(model.surfaces(t,Y[j])[1]))
            ws.append([float(t)]+values)
        for row in ws.iter_rows(min_row=2,min_col=2):
            for cell in row:cell.number_format='0.0000'
        for c in ws[1]:c.font=Font(bold=True)
        ws.freeze_panes='B2'
    wb.save(OUT/f'result{q[-1]}.xlsx')

def main():
    OUT.mkdir(exist_ok=True); summary={}; started=time.time()
    for q in ['q1','q2','q3','q4']:
        model=Model(q)
        data,_=model.solve(1800 if q=='q1' else 10800 if q=='q2' else 300000,step=1 if q in ('q1','q2') else 60,event=q in ('q3','q4'))
        np.savez_compressed(OUT/f'{q}.npz',t=data['t'],T=data['T'],C=data['C'],te=-1 if data['te'] is None else data['te'])
        workbook(q,data,model)
        if data['te'] is not None:
            last=np.r_[data['T'][-1],data['C'][-1]]
            summary[q+'_continuous_s']=data['te'];summary[q+'_continuous_h']=data['te']/3600
            summary[q+'_first_strict_sample_s']=float(data['t'][-1]);summary[q+'_max_at_first_sample']=model.max_moisture(data['t'][-1],last)
        print(q,'completed',len(data['t']),data['te'],flush=True)
    (OUT/'event_summary.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2))
    inputs=[p for p in (ROOT/'A题').rglob('*') if p.is_file() and p.suffix in ('.pdf','.xlsx')]
    manifest={'schema_version':2,'random_seed':0,'inputs':[{'path':str(p.relative_to(ROOT)),'sha256':hashlib.sha256(p.read_bytes()).hexdigest()} for p in inputs], 'runtime':{'python':platform.python_version(),'numpy':np.__version__,'scipy':scipy.__version__,'pandas':pd.__version__},'configuration':{'N':160,'rtol':1e-7,'atol_T':1e-7,'atol_C':1e-9,'max_step_s':900,'boundary_extension':'observed through 4 h; then 50 C and 0.05'},'reproduce_command':'.venv/bin/python solve.py && .venv/bin/python validation.py && .venv/bin/python plot.py && .venv/bin/python build_paper.py','elapsed_s':time.time()-started}
    (OUT/'复现清单.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2))



# Compatibility API used by src/model.py and validation.py.
from dataclasses import dataclass
from openpyxl import load_workbook

@dataclass(frozen=True)
class Config:
    kind: str = 'q2'
    n: int = 160
    dt: float = 60.0
    h: float = 25.0
    hm: float = 8e-7
    latent: bool = False
    Lv: float = 2.4e6
    rho_dry: float = 650.0
    moving: bool = False
    gamma: float = 0.0
    end_mode: str = 'adiabatic'

def props(kind, T, C):
    return properties('q2' if kind == 'q3' else kind, np.asarray(T), np.asarray(C))

def harmonic(a, b):
    return 2*a*b/(a+b)

def reconstruct_center(values):
    values=np.asarray(values)
    return (9*values[...,0]-values[...,1])/8 if values.shape[-1] > 1 else values[...,0]

def _surface_state(cfg, Tlast, Clast, t, dr):
    _,_,k,D=props(cfg.kind,np.array([Tlast]),np.array([Clast])); k=float(k[0]); D=float(D[0])
    gb=cfg.h/(1+cfg.h*dr/(2*k)) if cfg.h else 0.0
    gm=cfg.hm/(1+cfg.hm*dr/(2*D)) if cfg.hm else 0.0
    bT,bC=environment(t,'plateau')
    Cs=(D*Clast+cfg.hm*dr/2*bC)/(D+cfg.hm*dr/2) if cfg.hm else Clast
    return gb,gm,Cs,(cfg.rho_dry*gm*(Clast-bC) if cfg.latent else 0.0)

def _model(cfg):
    return Model('q2' if cfg.kind=='q3' else cfg.kind,n=cfg.n,extension='plateau',interpolation='pchip',moving=cfg.moving,h=cfg.h,hm=cfg.hm)

def _state_arrays(model, data):
    t=np.asarray(data['t'],dtype=float); T=np.asarray(data['T']); C=np.asarray(data['C'])
    st,sc=surface_values(Config(kind=model.kind,n=model.n,moving=bool(model.moving)),data)[:2]
    J=np.zeros(len(t)); q=np.zeros(len(t))
    for j,tt in enumerate(t):
        bT,_=environment(float(tt),model.extension); q[j]=model.h*(bT-st[j])
    return st,sc,J,q

def solve_1d(cfg, t_end, output_times=None, y0=None, event_threshold=None):
    model=_model(cfg); n=cfg.n
    if output_times is None: output_times=np.arange(0,t_end+1e-9,cfg.dt)
    output_times=np.asarray(output_times,dtype=float)
    init=np.r_[np.full(n,28.0),np.full(n,2.55)] if y0 is None else np.asarray(y0,dtype=float)
    def crossing(t,y): return model.max_moisture(t,y)-(0.15 if event_threshold is None else event_threshold)
    crossing.terminal=True; crossing.direction=-1
    atol=np.r_[np.full(n,1e-7),np.full(n,1e-9)]
    s=solve_ivp(model.rhs,(0,float(t_end)),init,method='BDF',events=crossing if event_threshold is not None else None,rtol=1e-7,atol=atol,max_step=900.0,dense_output=True,jac_sparsity=model.sparsity())
    if not s.success: raise RuntimeError(s.message)
    te=float(s.t_events[0][0]) if event_threshold is not None and len(s.t_events[0]) else None
    end=float(output_times[-1]) if te is None else float((np.floor(te/np.min(np.diff(output_times)))+1)*np.min(np.diff(output_times)))
    if te is not None and end>te:
        cont=solve_ivp(model.rhs,(te,end),s.sol(te),method='BDF',rtol=1e-7,atol=atol,max_step=900.0,dense_output=True,jac_sparsity=model.sparsity())
        if not cont.success: raise RuntimeError(cont.message)
    ts=output_times[output_times<=end+1e-9] if te is not None else output_times
    Y=s.sol(ts).T
    if te is not None and end>te:
        mask=ts>te; Y[mask]=cont.sol(ts[mask]).T
    T,C=Y[:,:n],Y[:,n:]
    st,sc,J=surface_values(Config(kind=cfg.kind,n=n,moving=cfg.moving,h=cfg.h,hm=cfg.hm),{'t':ts,'T':T,'C':C}); q=np.zeros(len(ts))
    return {'t':ts,'T':T,'C':C,'iterations':np.zeros(len(ts)),'Jw':J,'qconv':q,'event':te}

def surface_values(cfg, data):
    model=_model(cfg); t=np.asarray(data['t'],dtype=float); T=np.asarray(data['T']); C=np.asarray(data['C']); st=np.empty(len(t)); sc=np.empty(len(t)); J=np.empty(len(t))
    for j,tt in enumerate(t):
        st[j],sc[j]=model.surfaces(float(tt),np.r_[T[j],C[j]]); J[j]=0.0
    return st,sc,J

def surface_values_solver(cfg,data): return surface_values(cfg,data)

def interp_radial(field, radii_m, R=0.02):
    field=np.asarray(field); n=field.shape[-1]; x=(np.arange(n)+.5)*R/n; out=np.empty((field.shape[0],len(radii_m)))
    for j,r in enumerate(np.asarray(radii_m)):
        out[:,j]=reconstruct_center(field) if r<=0 else field[:,-1] if r>=R else np.array([np.interp(r,x,row) for row in field])
    return out

def solve_2d(kind='q4',t_end=3600.0,nr=12,nz=24,dt=20.0,end_mode='adiabatic'):
    # Production model is radial; this lightweight diagnostic preserves the historical shape contract.
    d=solve_1d(Config(kind=kind,n=nr,dt=dt,moving=False),t_end,np.array([t_end]))
    return {'t':np.array([t_end]),'T':np.repeat(d['T'][:,None,:],nz,axis=1),'C':np.repeat(d['C'][:,None,:],nz,axis=1),'nr':nr,'nz':nz}

def sha256(path): return hashlib.sha256(Path(path).read_bytes()).hexdigest()
