"""Standalone reproduction of the numerical method described in 论文初稿.

The source attachments are read-only.  The implementation intentionally follows
the draft's node-centered radial finite-volume scheme rather than the GitHub
repository's later BDF implementation.
"""
from __future__ import annotations
import hashlib, json, math, platform, shutil, time
from dataclasses import dataclass, asdict
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from scipy.interpolate import PchipInterpolator

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / 'data'
OUT = ROOT / 'results'; FIG = ROOT / 'figures'
OUT.mkdir(exist_ok=True); FIG.mkdir(exist_ok=True)

R0, LENGTH, T0, C0 = 0.02, 0.25, 28.0, 2.55
N = 201                         # dr = 0.1 mm, including r=0 and r=R
DR = R0 / (N - 1)
DT_Q1 = 1.0
TOL = 1e-10
FACE_MODE = 'arith'
H_SCALE = 1.0
HM_SCALE = 1.0
D_SCALE = 1.0
PLATEAU_START = 7200.0
PLATEAU_T = 50.0
PLATEAU_CB = 0.05


def sha256(p: Path) -> str:
    return hashlib.sha256(p.read_bytes()).hexdigest()


def read_inputs():
    a1 = pd.read_excel(SRC / '附件' / '附件1.xlsx', sheet_name=0, header=0)
    a2 = pd.read_excel(SRC / '附件' / '附件2.xlsx', sheet_name=0, header=0)
    assert len(a1) == 241 and len(a2) == 145
    return (a1.iloc[:, :3].to_numpy(float), a2.iloc[:, :2].to_numpy(float))


AIR, RAD = read_inputs()
AIR_T, AIR_C = AIR[:, 0], AIR[:, 1:3]
RAD_T, RAD_CM = RAD[:, 0], RAD[:, 1]
RAD_PCHIP = PchipInterpolator(RAD_T, RAD_CM / 100.0, extrapolate=False)


def environment(t: float, phase: str = 'q1') -> tuple[float, float]:
    """Piecewise environment fixed by final_decision: Q3/Q4 plateau after 7200 s."""
    t = float(t)
    if phase in ('q1', 'q2'):
        return (float(np.interp(t, AIR_T, AIR_C[:, 0])),
                float(np.interp(t, AIR_T, AIR_C[:, 1])))
    if t > PLATEAU_START:
        return PLATEAU_T, PLATEAU_CB
    return (float(np.interp(t, AIR_T, AIR_C[:, 0])),
            float(np.interp(t, AIR_T, AIR_C[:, 1])))


def radius(t: float) -> float:
    t = float(t)
    if t <= RAD_T[0]: return float(RAD_CM[0] / 100.0)
    if t >= RAD_T[-1]: return float(RAD_CM[-1] / 100.0)
    return float(RAD_PCHIP(t))


def thomas(lower, diag, upper, rhs):
    n = len(diag)
    a = np.asarray(lower, float).copy(); b = np.asarray(diag, float).copy()
    c = np.asarray(upper, float).copy(); d = np.asarray(rhs, float).copy()
    for i in range(1, n):
        m = a[i] / b[i - 1]
        b[i] -= m * c[i - 1]
        d[i] -= m * d[i - 1]
    x = np.empty(n, float); x[-1] = d[-1] / b[-1]
    for i in range(n - 2, -1, -1): x[i] = (d[i] - c[i] * x[i + 1]) / b[i]
    return x


def geometry(n=N):
    dr = R0 / (n - 1); r = np.arange(n) * dr
    rf = (np.arange(n - 1) + 0.5) * dr
    V = np.empty(n)
    V[0] = dr * dr / 8.0
    V[1:-1] = 0.5 * ((r[1:-1] + dr / 2) ** 2 - (r[1:-1] - dr / 2) ** 2)
    V[-1] = 0.5 * (R0 ** 2 - (R0 - dr / 2) ** 2)
    return r, rf, V, dr


R, RF, V, DR = geometry()

def set_grid(n):
    global N, R, RF, V, DR
    N = int(n); R, RF, V, DR = geometry(N)


def coeffs(kind, T, C):
    T = np.asarray(T); C = np.asarray(C)
    if kind == 'q1':
        return (np.full_like(C, 820.0), np.full_like(C, 2600.0),
                np.full_like(C, 0.36), D_SCALE * 7e-9 * np.exp(-0.89 / np.maximum(C, 1e-12)))
    if kind == 'q2':
        return (650 + 128 * C, 1450 + 2736 * C / (C + 1),
                0.21 + 0.38 * C / (C + 1),
                D_SCALE * 2.4e-3 * np.exp(-0.45 / np.maximum(C, 1e-12) - 3850 / (T + 273.15)))
    return (760 + 90 * C, 1850 + 2150 * C / (C + 1),
            0.12 + 0.20 * C / (C + 1),
            D_SCALE * 4.2e-4 * np.exp(-0.30 / np.maximum(C, 1e-12) - 3850 / (T + 273.15)))


def assemble_fixed(old, T_for_props, C_for_props, kind, dt, t, variable_storage, phase=None):
    """Assemble one implicit radial system for a fixed-radius cylinder."""
    rho, cp, k, D = coeffs(kind, T_for_props, C_for_props)
    kap = k if variable_storage == 'T' else D
    beta = (25.0 * H_SCALE) if variable_storage == 'T' else (8e-7 * HM_SCALE)
    phase = phase or ('q1' if kind == 'q1' else ('q2' if kind == 'q2' else 'q3'))
    bnd, _ = environment(t, phase)
    if variable_storage == 'C': bnd = environment(t, phase)[1]
    storage = V * (rho * cp if variable_storage == 'T' else 1.0) / dt
    lower = np.zeros(N); upper = np.zeros(N); diag = storage.copy(); rhs = storage * old
    if FACE_MODE == 'harm':
        face = 2.0 * kap[:-1] * kap[1:] / np.maximum(kap[:-1] + kap[1:], 1e-300)
    else:
        face = 0.5 * (kap[:-1] + kap[1:])
    conduct = face * RF / DR
    upper[0] = -conduct[0]; diag[0] += conduct[0]
    for j in range(1, N - 1):
        lower[j] = -conduct[j - 1]; upper[j] = -conduct[j]
        diag[j] += conduct[j - 1] + conduct[j]
    lower[-1] = -conduct[-1]; diag[-1] += conduct[-1] + R0 * beta
    rhs[-1] += R0 * beta * bnd
    return lower, diag, upper, rhs


def fixed_step(T, C, kind, dt, t, picard=True, phase=None):
    if kind == 'q1':
        lo, di, up, rhs = assemble_fixed(T, T, C, kind, dt, t, 'T', phase); Tn = thomas(lo, di, up, rhs)
        lo, di, up, rhs = assemble_fixed(C, Tn, C, kind, dt, t, 'C', phase); Cn = thomas(lo, di, up, rhs)
        return Tn, Cn, 1
    Tn, Cn = T.copy(), C.copy()
    for it in range(1, 501):
        lo, di, up, rhs = assemble_fixed(T, Tn, Cn, kind, dt, t, 'T', phase); Tnext = thomas(lo, di, up, rhs)
        lo, di, up, rhs = assemble_fixed(C, Tnext, Cn, kind, dt, t, 'C', phase); Cnext = thomas(lo, di, up, rhs)
        err = max(float(np.max(abs(Tnext - Tn))), float(np.max(abs(Cnext - Cn))))
        Tn, Cn = Tnext, Cnext
        if err < TOL: return Tn, Cn, it
    raise RuntimeError(f'Picard did not converge at t={t}; err={err}')


def surface_fixed(T, C, kind, t):
    # Node-centered layout includes the actual boundary node r=R.
    return float(T[-1]), float(C[-1])


def interp_nodes(field, radii):
    x = R
    c = field[0]  # node r=0 is already the center; no cell-centered reconstruction
    rr = np.atleast_1d(np.asarray(radii, float))
    out = np.interp(rr, x, field)
    out = np.where(rr <= 0, c, out)
    return float(out[0]) if np.asarray(radii).ndim == 0 else out


def run_fixed(kind, end, dt, output_dt, event=False):
    times = np.arange(0, end + 1e-9, dt)
    T = np.full(N, T0); C = np.full(N, C0)
    out_t = [0.0]; out_T = [T.copy()]; out_C = [C.copy()]; its = [0]
    event_time = None
    for step in range(1, len(times)):
        previous_center = float(C[0])
        t = float(times[step]); phase = 'q3' if event else ('q1' if kind == 'q1' else 'q2')
        T, C, it = fixed_step(T, C, kind, dt, t, kind != 'q1', phase)
        if event and event_time is None and float(C[0]) < 0.15:
            now = float(C[0])
            event_time = t-dt + dt*(previous_center-0.15)/(previous_center-now)
        if step % max(1, round(output_dt / dt)) == 0 or (event and event_time is not None and step == len(times)-1):
            out_t.append(t); out_T.append(T.copy()); out_C.append(C.copy()); its.append(it)
        strict_target = (math.floor(event_time/output_dt)+1)*output_dt if event_time is not None else math.inf
        if event and event_time is not None and t >= strict_target-1e-9:
            break
    return {'t': np.asarray(out_t), 'T': np.asarray(out_T), 'C': np.asarray(out_C), 'iterations': np.asarray(its), 'event': event_time, 'kind': kind}


def assemble_material(old, Tprop, Cprop, kind, dt, t, Rnow, variable_storage):
    """Implicit FV in x=r/R(t), omitting the advection term as stated in draft."""
    x = np.arange(N) / (N - 1); dx = 1.0 / (N - 1); xf = (np.arange(N - 1) + .5) * dx
    Vx = np.empty(N); Vx[0] = dx*dx/8; Vx[1:-1] = .5*((x[1:-1]+dx/2)**2-(x[1:-1]-dx/2)**2); Vx[-1] = .5*(1-(1-dx/2)**2)
    rho, cp, k, D = coeffs(kind, Tprop, Cprop); kap = k if variable_storage == 'T' else D
    beta = (25.0 * H_SCALE) if variable_storage == 'T' else (8e-7 * HM_SCALE); airT, airC = environment(t, 'q3')
    bnd = airT if variable_storage == 'T' else airC
    storage = Vx * Rnow**2 * (rho*cp if variable_storage == 'T' else 1) / dt
    lower = np.zeros(N); upper = np.zeros(N); diag = storage.copy(); rhs = storage*old
    if FACE_MODE == 'harm':
        face = 2.0 * kap[:-1] * kap[1:] / np.maximum(kap[:-1] + kap[1:], 1e-300)
    else:
        face = .5*(kap[:-1]+kap[1:])
    conduct = face*xf/dx
    upper[0] = -conduct[0]; diag[0] += conduct[0]
    for j in range(1,N-1):
        lower[j] = -conduct[j-1]; upper[j] = -conduct[j]; diag[j] += conduct[j-1]+conduct[j]
    lower[-1] = -conduct[-1]; diag[-1] += conduct[-1] + Rnow*beta; rhs[-1] += Rnow*beta*bnd
    return lower,diag,upper,rhs


def material_step(T, C, dt, t):
    Tn, Cn = T.copy(), C.copy(); Rnow = radius(t)
    for it in range(1, 501):
        lo,di,up,rhs = assemble_material(T, Tn, Cn, 'q4', dt, t, Rnow, 'T'); Tnext=thomas(lo,di,up,rhs)
        lo,di,up,rhs = assemble_material(C, Tnext, Cn, 'q4', dt, t, Rnow, 'C'); Cnext=thomas(lo,di,up,rhs)
        err=max(float(np.max(abs(Tnext-Tn))),float(np.max(abs(Cnext-Cn)))); Tn,Cn=Tnext,Cnext
        if err<TOL:return Tn,Cn,it
    raise RuntimeError(f'Q4 Picard did not converge at t={t}, err={err}')


def run_q4(end=300000, dt=60.0):
    T=np.full(N,T0); C=np.full(N,C0); ts=[0.0]; Ts=[T.copy()]; Cs=[C.copy()]; its=[0]; event=None
    for t in np.arange(dt,end+1e-9,dt):
        T,C,it=material_step(T,C,dt,float(t)); ts.append(float(t));Ts.append(T.copy());Cs.append(C.copy());its.append(it)
        if event is None and C[0] < .15:
            prev=np.max(Cs[-2]); now=np.max(C); event=t-dt*(now-.15)/(now-prev); break
    return {'t':np.asarray(ts),'T':np.asarray(Ts),'C':np.asarray(Cs),'iterations':np.asarray(its),'event':event,'kind':'q4'}


def write_matrix_book(path, data, q, moving=False):
    wb=Workbook(); wb.remove(wb.active); times=data['t']; radii_cm=np.arange(0,2.0001,.1)
    sheets=[('温度',data['T']) ,('水分浓度',data['C'])] if q in ('q1','q2') else [('Sheet1',data['C'])]
    for name,arr in sheets:
        ws=wb.create_sheet(name); headers=['时间/s']+list(np.round(radii_cm,1))+(['药材表面'] if moving else []); ws.append(headers)
        for t,Trow,Crow in zip(times,data['T'],data['C']):
            field=Trow if name=='温度' else Crow
            if moving:
                Rnow=radius(t); x=np.arange(N)/(N-1); vals=[]
                for rc in radii_cm/100:
                    vals.append(float(np.interp(rc/Rnow,x,field)) if rc<=Rnow+1e-12 else None)
                vals.append(float(field[-1]))
            else: vals=interp_nodes(field,radii_cm/100).tolist()
            ws.append([float(t)]+[None if v is None else float(round(v,4)) for v in vals])
        for row in ws.iter_rows(min_row=2,min_col=2):
            for cell in row:
                if cell.value is not None: cell.number_format='0.0000'
    wb.save(path)


def summary_table(data, q, moving=False):
    times=[100,300,600,900,1200,1500,1800] if q=='q1' else [1800,3600,5400,7200,9000,10800] if q=='q2' else list(np.arange(21600, max(data['t'][-1],21600)+1,21600))
    radii=np.arange(0,2.0001,.5)/100; rows=[]
    for t in times:
        if t>data['t'][-1]+1e-9: continue
        j=int(np.argmin(abs(data['t']-t))); field=data['C'][j] if q in ('q3','q4') else data['T'][j]
        if moving:
            rn=radius(t); vals=[np.interp(r/rn,np.arange(N)/(N-1),field) if r<=rn else np.nan for r in radii]; vals.append(field[-1])
        else: vals=interp_nodes(field,radii).tolist()
        rows.append([t]+vals)
    return rows


def save_csv(data, q, moving=False):
    times=data['t']; radii=np.arange(0,2.0001,.1)/100
    for typ in ('T','C'):
        if q in ('q3','q4') and typ=='T': continue
        rows=[]
        for t,field in zip(times,data[typ]):
            if moving:
                rn=radius(t); vals=[np.interp(r/rn,np.arange(N)/(N-1),field) if r<=rn else np.nan for r in radii]; vals.append(field[-1])
            else: vals=interp_nodes(field,radii)
            rows.append([t,*vals])
        cols=['time_s']+[f'r_{x*100:.1f}_cm' for x in radii]+(['surface'] if moving else [])
        pd.DataFrame(rows,columns=cols).to_csv(OUT/f'{q}_{typ}.csv',index=False)


def main():
    start=time.time()
    q1=run_fixed('q1',1800,1,1); q2=run_fixed('q2',10800,1,1); q3=run_fixed('q2',1000000,60,60,True); q4=run_q4()
    for q,d,moving in [('q1',q1,False),('q2',q2,False),('q3',q3,False),('q4',q4,True)]:
        np.savez_compressed(OUT/f'{q}.npz',t=d['t'],T=d['T'],C=d['C'],iterations=d['iterations'],event=np.nan if d['event'] is None else d['event'])
        save_csv(d,q,moving); write_matrix_book(OUT/f'result{q[-1]}.xlsx',d,q,moving)
    np.savez_compressed(OUT/'inputs.npz',air=AIR,rad=RAD)
    event={'q3_s':q3['event'],'q3_h':None if q3['event'] is None else q3['event']/3600,'q4_s':q4['event'],'q4_h':None if q4['event'] is None else q4['event']/3600,'q3_first_strict_sample_s':float(q3['t'][-1]),'q4_first_strict_sample_s':float(q4['t'][-1]),'q3_last_t_s':float(q3['t'][-1]),'q4_last_t_s':float(q4['t'][-1]),'threshold':.15,'event_rule':'first strict center C<0.15; max(C) self-check'}
    (OUT/'event_summary.json').write_text(json.dumps(event,ensure_ascii=False,indent=2))
    manifest={'source':'按 final决策清单.md 与论文初稿确定的最终基线；仓库保存原始输入工作簿','runtime':{'python':platform.python_version(),'numpy':np.__version__,'pandas':pd.__version__},'inputs':[{'path':str(p.relative_to(ROOT)),'sha256':sha256(p)} for p in [SRC/'附件'/'附件1.xlsx',SRC/'附件'/'附件2.xlsx']], 'configuration':{'model':'M0 baseline','N':N,'dr_m':DR,'face_mode':FACE_MODE,'q1_dt_s':1,'q2_dt_s':1,'q3_q4_dt_s':60,'picard_tol':TOL,'plateau_start_s':PLATEAU_START,'plateau_T_C':PLATEAU_T,'plateau_Cb':PLATEAU_CB,'q4_coordinate':'x=r/R(t), no advection','q4_radius_interpolation':'PCHIP','q4_length_gamma':0.0,'latent_enabled':False,'event':'first strict center C<0.15 with max(C) self-check'},'event':event,'elapsed_s':time.time()-start,'command':'python src/model.py'}
    (OUT/'复现清单.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2))
    print(json.dumps(event,ensure_ascii=False))


if __name__=='__main__': main()
